"""Word rendering of audit results; performs no configuration evaluation."""
from datetime import date
from pathlib import Path
from .catalog import STATUSES, STATUS_RGB, NOT_EVALUATED


def build_report(account_results, domain_sections_full, out_path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_heading("Cloudflare Configuration Audit", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"Automated report -- generated {date.today().isoformat()} from the repository's YAML configuration")
    run.italic = True

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Scope and limitations: ").bold = True
    p.add_run(
        "This report grades every control from the Cloudflare Account Audit and Cloudflare Zone Audit "
        "checklists that can be answered by reading the YAML in account/ and domains/*. It is generated "
        "by scripts/compliance_audit.py and is not a substitute for the live Cloudflare/GitHub checks, "
        "identity reviews and sign-off process in the source checklists. Controls this script cannot "
        "evaluate are listed in the appendix with the reason, not guessed at."
    )

    # ---- executive summary ----
    doc.add_heading("Executive summary", level=1)
    all_sections = [("Account", account_results)] + [(n, r) for n, r, _ in domain_sections_full]
    summary_table = doc.add_table(rows=1, cols=2 + len(STATUSES))
    summary_table.style = "Light Grid Accent 1"
    hdr = summary_table.rows[0].cells
    hdr[0].text = "Root"
    for i, s in enumerate(STATUSES):
        hdr[1 + i].text = s
    hdr[-1].text = "Total"
    for root_name, results in all_sections:
        row = summary_table.add_row().cells
        row[0].text = root_name
        counts = {s: sum(1 for r in results if r.status == s) for s in STATUSES}
        for i, s in enumerate(STATUSES):
            row[1 + i].text = str(counts[s])
        row[-1].text = str(len(results))

    doc.add_paragraph()
    doc.add_paragraph(
        f"{len(NOT_EVALUATED)} additional control(s) across both source checklists need a live Cloudflare/"
        "GitHub check, external verification, or a Terraform module this repo does not have -- see the "
        "appendix."
    )

    # ---- per-root detail ----
    def add_detail_table(results):
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [0.6, 1.3, 2.6, 0.7, 0.8, 3.0]
        for cell, w in zip(table.rows[0].cells, widths):
            cell.width = Inches(w)
        for i, h in enumerate(["ID", "Area", "Control", "Severity", "Result", "Evidence"]):
            table.rows[0].cells[i].text = h
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for r in results:
            cells = table.add_row().cells
            cells[0].text = r.id
            cells[1].text = r.area
            cells[2].text = r.control
            cells[3].text = r.severity
            run = cells[4].paragraphs[0].add_run(r.status)
            run.bold = True
            run.font.color.rgb = RGBColor(*STATUS_RGB[r.status])
            cells[5].text = r.evidence

    doc.add_heading("Account", level=1)
    add_detail_table(account_results)

    for root_name, results, is_example in domain_sections_full:
        h = doc.add_heading(root_name, level=1)
        if is_example:
            run = h.add_run("  (EXAMPLE DATA -- fake IDs, not a real zone)")
            run.italic = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        add_detail_table(results)

    # ---- appendix ----
    doc.add_heading("Appendix: controls not evaluated by this script", level=1)
    doc.add_paragraph(
        "These come from the same two source checklists but need a live Cloudflare/GitHub API check, "
        "external verification (registrar, origin server, OSINT), process evidence, or a Terraform "
        "resource this repository does not currently model."
    )
    app_table = doc.add_table(rows=1, cols=4)
    app_table.style = "Light Grid Accent 1"
    for i, h in enumerate(["ID", "Control", "Checklist", "Why this script can't grade it"]):
        app_table.rows[0].cells[i].text = h
        app_table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for cid, control, checklist, reason in NOT_EVALUATED:
        cells = app_table.add_row().cells
        cells[0].text = cid
        cells[1].text = control
        cells[2].text = checklist
        cells[3].text = reason

    # python-docx's default template omits w:percent on w:zoom, which fails
    # strict OOXML validation (Word itself tolerates it, but don't rely on that).
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    # python-docx stamps author="python-docx" / comments="generated by
    # python-docx" by default; clear that machine-generated fingerprint.
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
